"""Demo dataset: a FICTIONAL company's 24 months of monthly P&L data.

The demo is rendered into real PDF / XLSX / CSV files which then go through the
exact same extraction -> normalization -> validation -> calculation pipeline as a
user upload. Nothing downstream is special-cased for the demo.

The data deliberately contains realistic imperfections so the controls are visible:
  * a missing value ("Travel & Entertainment", Feb 2026)                -> REVIEW REQUIRED
  * a label the taxonomy cannot map ("Misc. adjustments (see note 7)")   -> REVIEW REQUIRED
  * a fuzzy label ("Depreciation & Amort.")                              -> WARNING (confirm)
  * a one-off rent spike (lease exit fee, Jun 2026)                      -> WARNING (outlier)
  * a typo in the printed Net Income subtotal (Mar 2026)                 -> WARNING (reconciliation)
  * two payroll rows that aggregate into one category                    -> lineage shows both
  * a balance-sheet memo row ("Trade Payables")                          -> mapped to Accounts Payable
"""
from __future__ import annotations

import io
from datetime import datetime

from app.extraction.parsing import MONTH_NAMES

COMPANY = "Meridian Ridge Components, Inc. (fictional)"
DEMO_BANNER = "DEMO DATA - Not for financial decision-making"
START = (2024, 9)  # 24 months: Sep 2024 .. Aug 2026
N_MONTHS = 24
SEASONALITY = [0.97, 1.01, 1.04, 1.08, 0.93, 0.95, 1.00, 1.02, 1.03, 0.99, 0.98, 1.00]

MISSING_CELL = ("Travel & Entertainment", "2026-02")
NET_INCOME_TYPO = "2026-03"
RENT_SPIKE = "2026-06"


def _periods() -> list[tuple[int, int]]:
    y, m = START
    out = []
    for _ in range(N_MONTHS):
        out.append((y, m))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def build_rows() -> tuple[list[str], list[tuple[str, list]]]:
    """Return (period keys, [(label, values)]) with integer dollar values."""
    periods = _periods()
    keys = [f"{y:04d}-{m:02d}" for y, m in periods]
    rows: dict[str, list] = {k: [] for k in [
        "Net Sales", "Cost of Sales", "Gross Profit", "Salaries & Wages", "Payroll Taxes & Benefits",
        "Rent & Occupancy", "Advertising & Promotion", "Software & IT", "Professional Fees",
        "Travel & Entertainment", "Misc. adjustments (see note 7)", "Total Operating Expenses", "EBITDA",
        "Depreciation & Amort.", "Interest Expense", "Income Tax Expense", "Net Income", "Trade Payables"]}
    for i, (y, m) in enumerate(periods):
        key = keys[i]
        growth = 1 + 0.0105 * i
        revenue = round(1_020_000 * growth * SEASONALITY[(m - 1) % 12], -2)
        cogs_rate = 0.585 - 0.0016 * i
        cogs = round(revenue * cogs_rate, -2)
        salaries = round(158_000 + 1_450 * i, -2)
        benefits = round(salaries * 0.21, -2)
        rent = 42_500 if i < 14 else 44_600
        if key == RENT_SPIKE:
            rent = 44_600 + 118_000   # one-off lease exit fee (real, but an outlier)
        marketing = round(revenue * (0.042 if m in (11, 12) else 0.031), -2)
        software = round(18_200 + 260 * i, -2)
        prof_fees = 14_500 if m % 3 else 27_800
        travel = round(6_100 + 90 * i + (2_400 if m in (3, 10) else 0), -2)
        misc = [1_150, -640, 2_310, 480, -920, 1_730][i % 6]
        opex = salaries + benefits + rent + marketing + software + prof_fees + travel + misc
        gross = revenue - cogs
        ebitda = gross - opex
        dep = 21_400 if i < 12 else 23_900
        interest = round(9_800 - 120 * i, -1)
        pretax = ebitda - dep - interest
        tax = round(max(pretax, 0) * 0.24, -2)
        net = pretax - tax
        net_printed = net
        if key == NET_INCOME_TYPO:
            s = str(int(net))
            net_printed = int(s[:-3] + s[-2] + s[-3] + s[-1])  # transposed digits in the printed subtotal
        payables = round(cogs * 0.62 + 35_000, -2)
        for label, v in [
            ("Net Sales", revenue), ("Cost of Sales", cogs), ("Gross Profit", gross),
            ("Salaries & Wages", salaries), ("Payroll Taxes & Benefits", benefits), ("Rent & Occupancy", rent),
            ("Advertising & Promotion", marketing), ("Software & IT", software),
            ("Professional Fees", prof_fees),
            ("Travel & Entertainment", None if key == MISSING_CELL[1] else travel),
            ("Misc. adjustments (see note 7)", misc), ("Total Operating Expenses", opex), ("EBITDA", ebitda),
            ("Depreciation & Amort.", dep), ("Interest Expense", interest), ("Income Tax Expense", tax),
            ("Net Income", net_printed), ("Trade Payables", payables),
        ]:
            rows[label].append(None if v is None else int(v))
    return keys, list(rows.items())


def _fmt(v) -> str:
    if v is None:
        return "-"
    return f"({abs(v):,})" if v < 0 else f"{v:,}"


def _label(key: str) -> str:
    y, m = key.split("-")
    return f"{MONTH_NAMES[int(m) - 1]} {y}"


# --------------------------------------------------------------------------- #
# Minimal PDF writer (text-only, standard Courier font) - no extra dependencies
# --------------------------------------------------------------------------- #
def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf() -> bytes:
    keys, rows = build_rows()
    per_page = 6
    pages_text: list[list[tuple[float, float, float, str]]] = []
    n_pages = (len(keys) + per_page - 1) // per_page
    for pi in range(n_pages):
        cols = keys[pi * per_page:(pi + 1) * per_page]
        lines: list[tuple[float, float, float, str]] = []
        lines.append((11, 40, 560, COMPANY))
        lines.append((9, 40, 544, "Monthly Income Statement (USD)"))
        lines.append((8, 40, 530, DEMO_BANNER))
        y = 500
        header = "Line Item".ljust(34) + "".join(_label(k).rjust(13) for k in cols)
        lines.append((8, 40, y, header))
        y -= 6
        lines.append((8, 40, y, "-" * len(header)))
        y -= 14
        for label, vals in rows:
            if label == "Trade Payables":
                y -= 8
                lines.append((8, 40, y, "Balance sheet memo"))
                y -= 14
            seg = vals[pi * per_page:(pi + 1) * per_page]
            lines.append((8, 40, y, label.ljust(34) + "".join(_fmt(v).rjust(13) for v in seg)))
            y -= 14
        lines.append((7, 40, 40, f"Page {pi + 1} of {n_pages}  |  Fictional company generated for product demonstration."))
        pages_text.append(lines)

    objects: list[bytes] = []

    def add(obj: bytes) -> int:
        objects.append(obj)
        return len(objects)

    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier /Encoding /WinAnsiEncoding >>")
    page_ids = []
    pages_id_placeholder = len(objects) + 1 + 2 * n_pages  # pages object added last
    for lines in pages_text:
        content = "\n".join(
            f"BT /F1 {size} Tf {x} {y} Td ({_pdf_escape(text)}) Tj ET" for size, x, y, text in lines
        ).encode("latin-1")
        cid = add(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        pid = add(("<< /Type /Page /Parent %d 0 R /MediaBox [0 0 792 612] /Resources << /Font << /F1 %d 0 R >> >> "
                   "/Contents %d 0 R >>" % (pages_id_placeholder, font_id, cid)).encode())
        page_ids.append(pid)
    pages_id = add(("<< /Type /Pages /Kids [%s] /Count %d >>" % (" ".join(f"{p} 0 R" for p in page_ids),
                                                                  len(page_ids))).encode())
    assert pages_id == pages_id_placeholder
    catalog_id = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)
    info_id = add(f"<< /Title ({_pdf_escape(COMPANY)} - Monthly P&L - DEMO) /Producer (Finance Intelligence Workbench) >>".encode("latin-1"))

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R /Info {info_id} 0 R >>\n"
              f"startxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


XLSX_LABELS = {
    "Net Sales": "Revenue from Operations", "Cost of Sales": "COGS", "Salaries & Wages": "Salaries",
    "Payroll Taxes & Benefits": "Employee Benefits", "Rent & Occupancy": "Rent Expense",
    "Advertising & Promotion": "Marketing", "Software & IT": "Software Subscriptions",
    "Trade Payables": "A/P",
}


def build_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    keys, rows = build_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly P&L"
    ws["A1"] = COMPANY
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = DEMO_BANNER
    ws["A3"] = "Amounts in USD"
    ws.cell(row=5, column=1, value="Line Item").font = Font(bold=True)
    for j, k in enumerate(keys, start=2):
        y, m = map(int, k.split("-"))
        c = ws.cell(row=5, column=j, value=datetime(y, m, 1))
        c.number_format = "mmm yyyy"
        c.font = Font(bold=True)
    for i, (label, vals) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=XLSX_LABELS.get(label, label))
        for j, v in enumerate(vals, start=2):
            if v is not None:
                ws.cell(row=i, column=j, value=v).number_format = "#,##0;(#,##0)"
    ws.column_dimensions["A"].width = 34
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Note 7: Miscellaneous adjustments include small accrual true-ups."
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv() -> bytes:
    keys, rows = build_rows()
    lines = ["Line Item," + ",".join(_label(k) for k in keys)]
    for label, vals in rows:
        lines.append('"' + label + '",' + ",".join("" if v is None else str(v) for v in vals))
    return ("\n".join(lines) + "\n").encode()


SAMPLES = {
    "meridian_ridge_monthly_pnl_DEMO.pdf": build_pdf,
    "meridian_ridge_monthly_pnl_DEMO.xlsx": build_xlsx,
    "meridian_ridge_monthly_pnl_DEMO.csv": build_csv,
}
