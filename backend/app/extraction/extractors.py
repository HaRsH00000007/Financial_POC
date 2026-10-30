"""Structural extraction of monthly financial tables from PDF, XLSX and CSV.

These extractors are rule-based: they locate period header rows and line-item
rows and record *where* every value was found (page / sheet / cell) together
with the raw token. They do not interpret meaning - mapping labels to financial
concepts is the normalization layer's job (optionally AI-assisted).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Optional

from app.extraction.parsing import (
    find_periods_in_line,
    is_number_token,
    parse_number,
    parse_period,
)
from app.models.domain import RawLineItem, SourceReference


@dataclass
class ExtractionResult:
    rows: list[RawLineItem]
    periods: list[str]
    pages: int | None = None
    sheets: list[str] = field(default_factory=list)
    scale: float = 1.0
    log: list[str] = field(default_factory=list)
    issues: list[dict] = field(default_factory=list)
    entity_name: str | None = None


class _RowBuilder:
    """Merges the same label across pages/sheets into one RawLineItem."""

    def __init__(self) -> None:
        self.rows: dict[str, RawLineItem] = {}
        self.order: list[str] = []

    def add(self, label: str, period: str, value: Optional[float], src: SourceReference, note: str | None = None):
        key = re.sub(r"\s+", " ", label.strip().lower())
        if key not in self.rows:
            rid = f"r{len(self.order) + 1:02d}"
            self.rows[key] = RawLineItem(id=rid, label=label.strip(), values={}, sources={}, confidence=1.0)
            self.order.append(key)
        row = self.rows[key]
        if period in row.values and row.values[period] is not None and value is not None \
                and row.values[period] != value:
            row.notes.append(f"Conflicting values for {period}: kept first ({row.values[period]}), "
                             f"ignored {value} from {src.sheet or 'page ' + str(src.page)}")
            return
        row.values[period] = value
        row.sources[period] = src
        row.confidence = min(row.confidence, src.confidence)
        if note:
            row.notes.append(note)

    def build(self) -> list[RawLineItem]:
        return [self.rows[k] for k in self.order]


def _scale_from_text(text: str) -> float:
    t = text.lower()
    if re.search(r"in\s+thousands|\(\$?000s?\)|usd\s*'?000", t):
        return 1000.0
    if re.search(r"in\s+millions|\(\$?m\)", t):
        return 1_000_000.0
    return 1.0


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def extract_pdf(data: bytes, doc_id: str, doc_name: str) -> ExtractionResult:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    builder = _RowBuilder()
    periods: set[str] = set()
    log: list[str] = []
    issues: list[dict] = []
    full_text = ""
    entity = None
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        full_text += text + "\n"
        header: list[str] = []
        rows_on_page = 0
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            if entity is None and page_no == 1 and not find_periods_in_line(line) and len(line) < 80:
                entity = line
            found = find_periods_in_line(line)
            if len(found) >= 2:
                header = found
                periods.update(found)
                continue
            if not header:
                continue
            tokens = line.split()
            nums: list[str] = []
            while tokens and is_number_token(tokens[-1]):
                nums.insert(0, tokens.pop())
            if len(nums) > len(header):
                # e.g. "Misc. adjustments (see note 7)" - surplus numeric-looking tokens belong to the label
                extra = len(nums) - len(header)
                tokens += nums[:extra]
                nums = nums[extra:]
            label = " ".join(tokens).strip(" .:")
            if not label or not nums or not re.search(r"[A-Za-z]", label):
                continue
            rows_on_page += 1
            if len(nums) != len(header):
                # Cannot align values to columns with certainty: record every period as
                # missing rather than guess which column each number belongs to.
                issues.append({"type": "column_mismatch", "page": page_no, "label": label,
                               "detail": f"Found {len(nums)} values for {len(header)} period columns; "
                                         "values not assigned."})
                for period in header:
                    builder.add(label, period, None, SourceReference(
                        document_id=doc_id, document_name=doc_name, page=page_no, row_label=label,
                        raw_text=" ".join(nums), extraction_method="pdf-text-table", confidence=0.0),
                        f"Page {page_no}: column count mismatch - values withheld")
                continue
            for period, tok in zip(header, nums):
                value, ok = parse_number(tok)
                conf = 0.0 if value is None else (0.96 if ok else 0.55)
                builder.add(label, period, value, SourceReference(
                    document_id=doc_id, document_name=doc_name, page=page_no, row_label=label,
                    raw_text=tok, extraction_method="pdf-text-table", confidence=conf))
        log.append(f"Page {page_no}: {len(header)} period columns, {rows_on_page} line-item rows")
    scale = _scale_from_text(full_text)
    return ExtractionResult(rows=builder.build(), periods=sorted(periods), pages=len(reader.pages),
                            scale=scale, log=log, issues=issues, entity_name=entity)


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #
def extract_xlsx(data: bytes, doc_id: str, doc_name: str) -> ExtractionResult:
    from openpyxl import load_workbook

    wb_values = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    wb_formulas = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    builder = _RowBuilder()
    periods: set[str] = set()
    log: list[str] = []
    issues: list[dict] = []
    text_blob = ""
    entity = None
    for ws in wb_values.worksheets:
        wf = wb_formulas[ws.title]
        header: dict[int, str] = {}
        rows_found = 0
        formula_cells = 0
        for row in ws.iter_rows():
            cells = list(row)
            parsed = {c.column: parse_period(c.value) for c in cells if c.value is not None}
            period_cols = {col: p for col, p in parsed.items() if p}
            if len(period_cols) >= 2:
                header = period_cols
                periods.update(period_cols.values())
                continue
            label_cell = next((c for c in cells if isinstance(c.value, str) and c.value.strip()), None)
            if label_cell is not None:
                text_blob += label_cell.value + "\n"
                if entity is None and not header:
                    entity = label_cell.value.strip()
            if not header or label_cell is None:
                continue
            if label_cell.column >= min(header):
                continue
            label = label_cell.value.strip()
            got_any = False
            for col, period in header.items():
                cell = ws.cell(row=label_cell.row, column=col)
                fcell = wf.cell(row=label_cell.row, column=col)
                formula = fcell.value if isinstance(fcell.value, str) and fcell.value.startswith("=") else None
                value, ok = parse_number(cell.value)
                conf = 0.99 if ok else 0.5
                note = None
                if formula and cell.value is None:
                    note = f"{cell.coordinate} contains formula {formula} with no cached value - treated as missing"
                    conf = 0.0
                if formula:
                    formula_cells += 1
                if value is None:
                    conf = 0.0
                if cell.value is not None or formula:
                    got_any = True
                builder.add(label, period, value, SourceReference(
                    document_id=doc_id, document_name=doc_name, sheet=ws.title, cell=cell.coordinate,
                    row_label=label, raw_text=formula or (None if cell.value is None else str(cell.value)),
                    extraction_method="xlsx-formula-cached-value" if formula else "xlsx-cell",
                    confidence=conf), note)
            if got_any:
                rows_found += 1
        log.append(f"Sheet '{ws.title}': {len(header)} period columns, {rows_found} line-item rows, "
                   f"{formula_cells} formula cells traced")
    return ExtractionResult(rows=builder.build(), periods=sorted(periods), sheets=wb_values.sheetnames,
                            scale=_scale_from_text(text_blob), log=log, issues=issues, entity_name=entity)


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #
def extract_csv(data: bytes, doc_id: str, doc_name: str) -> ExtractionResult:
    text = data.decode("utf-8-sig", errors="replace")
    reader = list(csv.reader(io.StringIO(text)))
    builder = _RowBuilder()
    header: dict[int, str] = {}
    periods: set[str] = set()
    rows_found = 0
    for r_idx, row in enumerate(reader, start=1):
        period_cols = {i: parse_period(v) for i, v in enumerate(row) if parse_period(v)}
        if len(period_cols) >= 2:
            header = period_cols
            periods.update(period_cols.values())
            continue
        if not header or not row or not row[0].strip():
            continue
        label = row[0].strip()
        rows_found += 1
        for col, period in header.items():
            tok = row[col] if col < len(row) else ""
            value, ok = parse_number(tok)
            builder.add(label, period, value, SourceReference(
                document_id=doc_id, document_name=doc_name, cell=f"R{r_idx}C{col + 1}", row_label=label,
                raw_text=tok, extraction_method="csv-cell",
                confidence=0.0 if value is None else (0.99 if ok else 0.5)))
    return ExtractionResult(rows=builder.build(), periods=sorted(periods), scale=_scale_from_text(text),
                            log=[f"CSV: {len(header)} period columns, {rows_found} line-item rows"])


def extract(data: bytes, file_type: str, doc_id: str, doc_name: str) -> ExtractionResult:
    if file_type == "pdf":
        return extract_pdf(data, doc_id, doc_name)
    if file_type == "xlsx":
        return extract_xlsx(data, doc_id, doc_name)
    if file_type == "csv":
        return extract_csv(data, doc_id, doc_name)
    raise ValueError(f"Unsupported file type: {file_type}")
